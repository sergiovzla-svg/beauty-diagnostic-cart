"""
Módulo de Exportación del Catálogo Excel a la Web App Estática
Lee 'nuevo_catalogo_final_consolidado.xlsx' y genera 'web/data/productos.json'
Copia las imágenes necesarias a 'web/images/' para garantizar que la web sea 100% autónoma.
"""
import argparse
import json
import shutil
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "data" / "nuevo_catalogo_final_consolidado.xlsx"
OUTPUT_IMAGES_DIR = BASE_DIR / "output_images"
WEB_DIR = BASE_DIR / "web"
WEB_DATA_DIR = WEB_DIR / "data"
WEB_IMAGES_DIR = WEB_DIR / "images"

# 5 Productos piloto para el lanzamiento inicial (opcional)
PILOT_IDS = {"LOR-0004", "LOR-0003", "LOR-0005", "LOR-0021", "LOR-0068"}

def exportar_catalogo_a_web(solo_piloto: bool = False):
    print("=== EXPORTANDO CATÁLOGO A LA WEB APP ESTÁTICA ===")
    print(f"Modo: {'PILOTO (5 productos)' if solo_piloto else 'COMPLETO (todos los productos activos)'}")
    
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"No se encontró el archivo Excel en {EXCEL_PATH}")

    WEB_DATA_DIR.mkdir(parents=True, exist_ok=True)
    WEB_IMAGES_DIR.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Control de Inventario"]

    productos = []
    omitidos = []
    
    for r in range(8, ws.max_row + 1):
        id_prod = ws.cell(r, 1).value
        if not id_prod or id_prod == "TOTALES GENERALES":
            continue
            
        id_str = str(id_prod).strip()
        nombre_comercial = str(ws.cell(r, 4).value or "").strip()

        # Filtrar productos marcados para eliminación en Columna A
        if "eliminar" in id_str.lower() or not id_str.upper().startswith("LOR-"):
            omitidos.append((r, id_str, nombre_comercial))
            continue

        if solo_piloto and id_str not in PILOT_IDS:
            continue

        formato = str(ws.cell(r, 5).value or "").strip()
        stock = int(ws.cell(r, 7).value or 0)

        # Precio Mercado (Col P / 16) y Mi Precio Venta (Col Q / 17)
        raw_mercado = ws.cell(r, 16).value
        try:
            p_mercado = int(float(str(raw_mercado).replace(',', '.'))) if raw_mercado is not None else 0
        except (ValueError, TypeError):
            p_mercado = 0

        raw_venta = ws.cell(r, 17).value
        try:
            p_venta = int(float(str(raw_venta).replace(',', '.'))) if raw_venta is not None else 0
        except (ValueError, TypeError):
            p_venta = 0

        foto = str(ws.cell(r, 20).value or "").strip()
        resena = str(ws.cell(r, 21).value or "").strip()
        rango = str(ws.cell(r, 22).value or "").strip()
        fuentes = str(ws.cell(r, 24).value or "").strip()
        titulo_ml = str(ws.cell(r, 26).value or "").strip()
        patron = str(ws.cell(r, 27).value or "Todos").strip()
        grosor = str(ws.cell(r, 28).value or "Todos").strip()
        raiz = str(ws.cell(r, 29).value or "Normal, Todos").strip()
        quimico = str(ws.cell(r, 30).value or "Natural, Todos").strip()
        necesidad = str(ws.cell(r, 31).value or "Cuidado General").strip()
        paso = str(ws.cell(r, 32).value or "1. Shampoo").strip()

        # Determinar nombre de imagen (preferir formato .webp optimizado)
        foto_final = ""
        if foto:
            foto_stem = Path(foto).stem
            foto_webp = f"{foto_stem}.webp"
            if (OUTPUT_IMAGES_DIR / foto_webp).exists():
                foto_archivo = foto_webp
            else:
                foto_archivo = foto

            src_img = OUTPUT_IMAGES_DIR / foto_archivo
            dst_img = WEB_IMAGES_DIR / foto_archivo
            if src_img.exists():
                shutil.copy2(src_img, dst_img)
            foto_final = f"images/{foto_archivo}"

        # Construir registro limpio para JavaScript
        item = {
            "id": id_str,
            "linea": str(ws.cell(r, 3).value or "").strip(),
            "nombre": nombre_comercial,
            "titulo_comercial": titulo_ml if titulo_ml else nombre_comercial,
            "formato": formato,
            "stock": stock,
            "precio_venta": p_venta,
            "precio_mercado": p_mercado,
            "foto": foto_final,
            "resena": resena,
            "rango_mercado": rango,
            "patron": [p.strip() for p in patron.split(",")],
            "grosor": [g.strip() for g in grosor.split(",")],
            "raiz": [rz.strip() for rz in raiz.split(",")],
            "quimico": [q.strip() for q in quimico.split(",")],
            "necesidad": necesidad,
            "paso": paso
        }
        productos.append(item)

    # Ordenar por paso de rutina
    productos.sort(key=lambda x: x["paso"])

    out_json = WEB_DATA_DIR / "productos.json"
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(productos, f, ensure_ascii=False, indent=2)

    out_js = WEB_DATA_DIR / "productos.js"
    with open(out_js, "w", encoding="utf-8") as f:
        f.write("// Catálogo generado automáticamente desde el inventario Excel\n")
        f.write("window.CATALOGO_PRODUCTOS = ")
        json.dump(productos, f, ensure_ascii=False, indent=2)
        f.write(";\n")

    print(f"Omitidos {len(omitidos)} productos marcados para eliminación:")
    for fila, id_txt, nom in omitidos:
        print(f"  - Fila {fila}: [{id_txt}] {nom}")

    print(f"\nExportados {len(productos)} productos activos exitosamente a: {out_json} y {out_js}")
    print(f"Imágenes listas en: {WEB_IMAGES_DIR}")
    return len(productos)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Exportar catálogo Excel a Web App")
    parser.add_argument("--pilot", action="store_true", help="Exportar solo piloto de 5 productos")
    parser.add_argument("--all", action="store_true", help="Exportar catálogo completo de productos activos (por defecto)")
    args = parser.parse_args()
    exportar_catalogo_a_web(solo_piloto=args.pilot)
